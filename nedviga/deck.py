"""v0.5.13 F-33: генератор презентаций (python-pptx) из подборок nedviga.

Шаблон нейтральный встроенный (мастер-слиды дефолтные), никаких внешних
файлов — продукт отдаёт структуру и данные, «дизайн» = чистая таблица.
Последний слайд — обязательный дисклеймер внешних источников.
"""
from __future__ import annotations

import statistics
from datetime import datetime, timezone
from pathlib import Path

DISCLAIMER = ("Данные получены из внешних источников (ЦИАН, Авито, open data Москвы) "
              "и могут быть неактуальны. Проверяйте цену, статус объявления и юридические "
              "параметры на источнике перед решениями. Прогноз пробок не включён: "
              "текущая загруженность ≠ прогноз час-пика.")

_CAP = 15


def _fmt_money(v) -> str:
    try:
        return f"{int(v):,}".replace(",", " ")
    except (TypeError, ValueError):
        return "—"


def _card_cells(c: dict) -> list[str]:
    pr = c.get("pricing") or {}
    pa = c.get("parameters") or {}
    bu = c.get("building") or {}
    loc = c.get("location") or {}
    return [
        str(c.get("id") or "—"),
        _fmt_money(pr.get("price_total_rub")),
        _fmt_money(pr.get("price_per_m2_rub")),
        f"{pa.get('rooms', '—')}/{pa.get('area_total_m2', '—')}/{pa.get('floor', '—')}-{pa.get('floors_total', '—')}",
        (bu.get("address_raw") or loc.get("address") or c.get("title") or "—")[:60],
    ]


def deck_from_cards(cards: list[dict], out_path: str, title: str = "Подборка недвижимости") -> dict:
    """Собирает pptx: титул → таблица топ-N → сводка рынка → дисклеймер
    (+ family-слайд если карточка семейной карты присутствует)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    cards = cards[:_CAP]
    prs = Presentation()

    # 1. титул
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    sub = s.placeholders[1] if len(s.placeholders) > 1 else None
    if sub is not None:
        sub.text = f"nedviga-cli · {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} · {len(cards)} объектов"

    # 2. таблица карточек
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Объекты"
    rows, cols = len(cards) + 1, 5
    tbl = s.shapes.add_table(rows, cols, Inches(0.3), Inches(1.4),
                             Inches(9.4), Inches(0.4 * rows)).table
    for j, h in enumerate(("ID", "Цена ₽", "₽/м²", "комн/м²/эт", "Адрес")):
        cell = tbl.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(11)
                r.font.bold = True
    for i, c in enumerate(cards, start=1):
        for j, v in enumerate(_card_cells(c)):
            cell = tbl.cell(i, j)
            cell.text = v
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

    # 3. family-слайд если есть
    fam = next((c for c in cards if c.get("family_card")), None)
    if fam is not None:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Семейная карта"
        body = s.placeholders[1] if len(s.placeholders) > 1 else None
        if body is not None:
            body.text = str(fam.get("verdict") or fam.get("summary") or "")[:800]

    # 4. сводка рынка (все числа с n — продуктовая константа)
    prices = [c["pricing"]["price_total_rub"] for c in cards
              if (c.get("pricing") or {}).get("price_total_rub")]
    ppms = [c["pricing"]["price_per_m2_rub"] for c in cards
            if (c.get("pricing") or {}).get("price_per_m2_rub")]
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Сводка"
    body = s.placeholders[1] if len(s.placeholders) > 1 else None
    lines = [f"Объектов в подборке: n={len(cards)}"]
    if prices:
        lines.append(f"Цена: медиана {_fmt_money(statistics.median(prices))} ₽, "
                     f"разброс {_fmt_money(min(prices))}–{_fmt_money(max(prices))} ₽ (n={len(prices)})")
    if ppms:
        lines.append(f"₽/м²: медиана {_fmt_money(statistics.median(ppms))} (n={len(ppms)})")
    if body is not None:
        body.text = "\n".join(lines)

    # 5. обязательный дисклеймер
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Ограничения"
    body = s.placeholders[1] if len(s.placeholders) > 1 else None
    if body is not None:
        body.text = DISCLAIMER

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return {"status": "ok", "slides": len(prs.slides), "cards": len(cards), "out": str(out)}
